from __future__ import annotations

import io
import os
from dataclasses import dataclass
from pathlib import Path

import pytesseract
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
import pypdfium2 as pdfium


def configure_tesseract() -> None:
    configured = os.getenv("CHTEST_TESSERACT_CMD")
    default_path = Path("D:/tools/tesseract/tesseract.exe")
    command = Path(configured) if configured else default_path
    if command.exists():
        pytesseract.pytesseract.tesseract_cmd = str(command)


configure_tesseract()


class DocumentExtractionError(Exception):
    pass


class DocumentExtractorUnavailableError(DocumentExtractionError):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    parser_name: str
    parser_version: str
    unit_count: int
    metadata: dict[str, object]


def extract_document_text(
    *,
    artifact_type: str,
    content: bytes,
    ocr_language: str = "eng+chi_sim",
) -> ExtractedDocument:
    if artifact_type == "context_pdf":
        return extract_pdf(content, ocr_language=ocr_language)
    if artifact_type == "context_xlsx":
        return extract_xlsx(content)
    if artifact_type == "context_image":
        return extract_image_ocr(content, language=ocr_language)
    raise DocumentExtractionError("Unsupported document artifact type.")


def extract_pdf(content: bytes, *, ocr_language: str = "eng+chi_sim") -> ExtractedDocument:
    pdfium_document = None
    ocr_page_count = 0
    ocr_failed_pages: list[int] = []
    effective_ocr_language = resolve_ocr_language(ocr_language)
    try:
        reader = PdfReader(io.BytesIO(content))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            page_has_images = bool(getattr(page, "images", []))
            ocr_text = ""
            if not page_text or page_has_images:
                try:
                    if pdfium_document is None:
                        pdfium_document = pdfium.PdfDocument(content)
                    rendered_page = pdfium_document[index - 1]
                    bitmap = rendered_page.render(scale=2.0)
                    try:
                        ocr_text = pytesseract.image_to_string(bitmap.to_pil(), lang=effective_ocr_language).strip()
                    finally:
                        bitmap.close()
                        rendered_page.close()
                    if ocr_text:
                        ocr_page_count += 1
                except (pytesseract.TesseractNotFoundError, pytesseract.TesseractError):
                    ocr_failed_pages.append(index)
                except Exception:
                    ocr_failed_pages.append(index)

            merged_text = merge_pdf_page_text(page_text, ocr_text)
            if merged_text:
                source_label = "OCR" if not page_text and ocr_text else "文字层 + OCR" if page_text and ocr_text else "文字层"
                pages.append(f"# Page {index} ({source_label})\n\n{merged_text}")
    except Exception as exc:
        raise DocumentExtractionError("PDF content could not be parsed.") from exc
    finally:
        if pdfium_document is not None:
            pdfium_document.close()

    text = "\n\n".join(pages).strip()
    if not text:
        if ocr_failed_pages:
            raise DocumentExtractorUnavailableError(
                f"PDF pages need OCR, but OCR is unavailable for page(s): {', '.join(map(str, ocr_failed_pages))}."
            )
        raise DocumentExtractionError("PDF contains no extractable text. Use image OCR for scanned pages.")

    metadata: dict[str, object] = {
        "page_count": len(reader.pages),
        "text_page_count": sum(1 for page in reader.pages if (page.extract_text() or "").strip()),
        "ocr_page_count": ocr_page_count,
        "ocr_language": effective_ocr_language if ocr_page_count else None,
    }
    if ocr_failed_pages:
        metadata["ocr_failed_pages"] = ocr_failed_pages
    return ExtractedDocument(
        text=text,
        parser_name="pypdf+tesseract" if ocr_page_count else "pypdf",
        parser_version="v2",
        unit_count=len(reader.pages),
        metadata=metadata,
    )


def merge_pdf_page_text(text_layer: str, ocr_text: str) -> str:
    text_layer = " ".join(text_layer.split())
    ocr_text = " ".join(ocr_text.split())
    if not text_layer:
        return ocr_text
    if not ocr_text:
        return text_layer
    normalized_text = text_layer.casefold()
    normalized_ocr = ocr_text.casefold()
    if normalized_ocr in normalized_text:
        return text_layer
    if normalized_text in normalized_ocr:
        return ocr_text
    return f"{text_layer}\n\n[图片 OCR]\n{ocr_text}"


def resolve_ocr_language(requested: str) -> str:
    requested_languages = [item.strip() for item in requested.split("+") if item.strip()]
    try:
        available_languages = set(pytesseract.get_languages(config=""))
    except (pytesseract.TesseractNotFoundError, pytesseract.TesseractError) as exc:
        raise DocumentExtractorUnavailableError(
            "Tesseract OCR is not installed or is not available on PATH."
        ) from exc

    selected_languages = [item for item in requested_languages if item in available_languages]
    if not selected_languages and "eng" in available_languages:
        selected_languages = ["eng"]
    if not selected_languages:
        raise DocumentExtractorUnavailableError(
            f"Tesseract does not provide any requested OCR language: {requested}."
        )
    return "+".join(selected_languages)


def extract_xlsx(content: bytes) -> ExtractedDocument:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sections: list[str] = []
        row_count = 0
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row_index, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
                normalized = ["" if value is None else str(value) for value in values]
                if not any(normalized):
                    continue
                rows.append(f"{row_index}\t" + "\t".join(normalized))
                row_count += 1
            if rows:
                sections.append(f"# Sheet: {worksheet.title}\n\n" + "\n".join(rows))
        workbook.close()
    except Exception as exc:
        raise DocumentExtractionError("XLSX content could not be parsed.") from exc
    text = "\n\n".join(sections).strip()
    if not text:
        raise DocumentExtractionError("XLSX contains no non-empty cells.")
    return ExtractedDocument(
        text=text,
        parser_name="openpyxl",
        parser_version="v1",
        unit_count=row_count,
        metadata={"sheet_count": len(sections), "row_count": row_count},
    )


def extract_image_ocr(content: bytes, *, language: str) -> ExtractedDocument:
    effective_language = resolve_ocr_language(language)
    try:
        with Image.open(io.BytesIO(content)) as image:
            image.load()
            width, height = image.size
            text = pytesseract.image_to_string(image, lang=effective_language).strip()
    except pytesseract.TesseractNotFoundError as exc:
        raise DocumentExtractorUnavailableError(
            "Tesseract OCR is not installed or is not available on PATH."
        ) from exc
    except pytesseract.TesseractError as exc:
        raise DocumentExtractionError("Tesseract OCR failed to process the image.") from exc
    except Exception as exc:
        raise DocumentExtractionError("Image content could not be parsed.") from exc
    if not text:
        raise DocumentExtractionError("OCR produced no text for the image.")
    return ExtractedDocument(
        text=f"# OCR Image\n\n{text}",
        parser_name="tesseract",
        parser_version="v1",
        unit_count=1,
        metadata={"ocr_language": effective_language, "width": width, "height": height},
    )
